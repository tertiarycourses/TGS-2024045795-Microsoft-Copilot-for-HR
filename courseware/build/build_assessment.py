#!/usr/bin/env python3
"""Generate aligned WSQ assessment papers and confidential answer keys."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from course_data import META, ASSESSMENT_DIR
from build_all import style_doc, header_footer, table, shade, convert_pdf, normalize_docx

VER="3.2"
LMS_URL="https://lms-tms.tertiaryinfotech.com/"

def assessment_header_footer(doc,label):
    for sec in doc.sections:
        p=sec.header.paragraphs[0]
        p.text=f"Tertiary Infotech Academy  |  {META['code']}  |  {META['title']}"
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{label} • Version {VER} • ")
        fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)

WA=[
 ("K1","Compare conventional HR process automation, Microsoft 365 Copilot and an agentic AI system. For each, explain the role of the human operator and give one suitable HR example.",
  ["Conventional automation follows pre-defined rules and is suitable for deterministic work such as routing a completed form.","Microsoft 365 Copilot generates, summarises or analyses while a person directs and verifies the work; an example is drafting a job description or summarising onboarding evidence.","An agent pursues a bounded goal by reasoning, retrieving knowledge and using approved tools; an example is a leave agent that gathers intent and invokes a governed approval flow.","Humans remain accountable, especially for consequential employment decisions and exceptions.","Course evidence: Unit 1, HR AI continuum, capability choice and operating-model slides."],20),
 ("K2","Explain four ways generative AI can improve HR work across the employee lifecycle. For each way, state one verification that an HR practitioner should perform before using the output.",
  ["Accept any four relevant lifecycle uses such as workforce planning, recruitment content, resume evidence extraction, policy drafting, onboarding, performance preparation, learning or employee communication.","Each use must include a concrete human verification such as checking source evidence, job-related criteria, policy version, factual accuracy, tone, privacy or approval.","The answer distinguishes acceleration from final decision authority.","Course evidence: Unit 2 and Labs 1–4, 6–7 and 10."],20),
 ("K3","Describe the six design elements of a reliable HR agent: purpose, knowledge, instructions, tools, boundaries and tests. Explain how the elements work together in a policy Q&A agent.",
  ["Purpose defines the user outcome and audience.","Knowledge contains approved, current policy and preserves permissions.","Instructions define response, citation, clarification, refusal and escalation behaviour.","Tools perform approved retrieval or transactions using least privilege.","Boundaries prohibit personalised entitlement or consequential decisions outside scope.","Tests cover normal, ambiguous, adversarial, permission and failure scenarios; together they make the agent useful, controlled and reviewable.","Course evidence: Unit 3 and Lab 7."],20),
 ("K4","Explain four responsible-AI controls that are especially important in HR and show how each control reduces a specific people or organisational risk.",
  ["Accept controls including purpose limitation/data minimisation, fair job-related criteria, transparency and citation, human oversight, least privilege, audit logging, security testing, escalation or monitoring.","Each control must be connected to a plausible HR risk such as discrimination, privacy breach, opaque decision, unauthorised action, misinformation or inability to contest.","A list without risk linkage is incomplete.","Course evidence: Unit 4 responsible-adoption and governance slides."],20),
 ("K5","Outline a 90-day approach for piloting an HR AI use case. Include the baseline, success and guardrail measures, human ownership, testing and scale/stop decision.",
  ["Days 1–30: define user outcome, baseline, scope/non-goals, risk/data review, synthetic-data prototype and owners.","Days 31–60: normal/edge/adversarial testing, human approvals, user feedback, operational design and remediation.","Days 61–90: measured pilot, incident rehearsal, value/quality/trust/risk review and formal scale/remediate/pause/retire decision.","Measures include at least one value, quality and risk guardrail; named owner and stop condition are explicit.","Course evidence: Unit 4, balanced scorecard and 90-day pathway slides."],20),
]

CASE_SCENARIO="""HarbourLight Logistics SG employs 620 people. Its five-person HR team is experiencing: 280 monthly policy questions; inconsistent hybrid-work answers; leave requests rekeyed from Teams messages; delayed onboarding access; and a hiring surge for a new distribution hub. Leaders propose Microsoft 365 Copilot for drafting and analysis, plus Copilot Studio agents for policy, onboarding and leave. A pilot dataset shows cycle time improved by 42%, but 8% of policy answers cited an obsolete document, two leave requests were duplicated after a connector retry, and candidate-screening summaries sometimes treated missing evidence as a negative score. Employee feedback requests clearer disclosure and a human appeal channel. The CHRO requires a safe 90-day recovery and scale decision."""

CS=[
 ("A1","Using the capability choices practised in Labs 5–9 and 11, classify the proposed HR uses as deterministic automation, Microsoft 365 Copilot-assisted work, Agent Builder knowledge work or Copilot Studio orchestration. Justify the lightest sufficient capability for at least four uses.",
  ["Classifications are plausible and clearly justified by task predictability, knowledge needs, integrations, consequence and human role.","Examples: drafting/summarising with Microsoft 365 Copilot; focused policy Q&A with Agent Builder; multi-step leave/onboarding with Copilot Studio plus deterministic flow; fixed validation or routing with automation.","The answer avoids using an agent where a simpler, safer tool is sufficient.","Lab trace: Labs 5–9 and 11."],20),
 ("A2","Apply the Lab 3 evidence-matrix pattern to design a fair, human-reviewed resume-screening workflow for the hiring surge. Include the rubric, evidence extraction, missing information, protected attributes, review, candidate communication and audit evidence.",
  ["Criteria are job-related and anchored before reviewing candidates.","Copilot extracts source-linked evidence and marks gaps as unknown rather than negative.","Protected/non-job-related attributes and proxies are excluded; human reviews all shortlist decisions.","The workflow includes consistency check, documented rationale, contestability/correction and confidential communication.","Inputs, prompts, extracted evidence, verification, revisions and final human decision are logged.","Lab trace: Lab 3."],20),
 ("A3","Apply Lab 8 to redesign the leave-agent workflow to prevent duplicate or unauthorised transactions. Show the conversation, deterministic validations, confirmation, approval, idempotency, failure handling and audit trail.",
  ["Agent gathers and clarifies intent; deterministic flow validates balance, dates, overlaps and approver.","The employee confirms before submission and an authorised manager approves exceptions.","A unique request/idempotency key and duplicate check prevent repeated transactions after retry.","Connector failure produces a safe status, no false success, retry/escalation path and compensating action where needed.","Audit records identity, inputs, decision, approval, tool result, error and notification.","Lab trace: Lab 8."],20),
 ("A4","Apply Lab 7 to propose immediate containment and remediation for the obsolete policy citations. Include knowledge governance, user communication, testing and ownership.",
  ["Pause or restrict affected answers/actions, remove or de-prioritise obsolete content and route uncertain cases to HR.","Identify authoritative owner, effective date, version and superseded records; preserve permissions.","Notify affected users where necessary and provide correction/human route.","Regression-test normal, ambiguous, outdated, adversarial and permission cases before reopening.","Name business, policy and technical owners plus monitoring and change control.","Lab trace: Lab 7."],20),
 ("A5","Use the Lab 12 control-room pattern to make a scale, remediate, pause or retire recommendation for each proposed capability and present a 90-day roadmap with balanced measures and human governance.",
  ["Recommendation uses the case evidence rather than cycle time alone.","Each capability receives a defensible decision; unsafe screening and stale/duplicate failures are remediated or paused until thresholds pass.","Roadmap has 30/60/90-day gates, accountable owners, employee feedback, incident rehearsal and retirement/rollback.","Measures cover service/value, accuracy/evidence, adoption/trust, override/escalation and control incidents.","Consequential employment decisions remain human and employees have disclosure, correction and appeal routes.","Lab trace: Lab 12."],20),
]

def cover(doc,title,key=False):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("TERTIARY INFOTECH ACADEMY PTE LTD\n").bold=True
    r=p.add_run("\n"+title.upper()+"\n"); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor(12,35,64)
    p.add_run(f"\n{META['title']}\nCourse Code: {META['code']}\n\nVersion {VER} • {META['date']}")
    if key:
        r=p.add_run("\n\nCONFIDENTIAL — ASSESSOR / TRAINER ONLY"); r.bold=True; r.font.color.rgb=RGBColor(196,54,62)
    doc.add_page_break()

def add_hyperlink(paragraph,url,text):
    r_id=paragraph.part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    link=OxmlElement("w:hyperlink"); link.set(qn("r:id"),r_id)
    run=OxmlElement("w:r"); rpr=OxmlElement("w:rPr")
    color=OxmlElement("w:color"); color.set(qn("w:val"),"0563C1"); rpr.append(color)
    underline=OxmlElement("w:u"); underline.set(qn("w:val"),"single"); rpr.append(underline)
    run.append(rpr); node=OxmlElement("w:t"); node.text=text; run.append(node); link.append(run); paragraph._p.append(link)

def answer_box(doc,answers=None,height_pt=86):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell=t.cell(0,0); cell.text=""
    if answers:
        p=cell.paragraphs[0]; r=p.add_run("Suggestive answers (not exhaustive):"); r.bold=True
        for item in answers: cell.add_paragraph("• "+item)
    else:
        trpr=t.rows[0]._tr.get_or_add_trPr(); trh=OxmlElement('w:trHeight'); trh.set(qn('w:val'),str(int(height_pt*20))); trh.set(qn('w:hRule'),'atLeast'); trpr.append(trh)
    doc.add_paragraph()

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def instructions(doc,title,competency):
    doc.add_heading("Trainee Information",1)
    table(doc,["Field","Candidate entry"],[("Trainee name (as per NRIC)",""),("Last 3 digits and alphabet of NRIC/FIN",""),("Assessment date",""),("Start / end time"," / ")])
    doc.add_heading("Instructions to Candidate",1)
    items=[
        "This is an individual, open-book assessment. Do not discuss answers with another person.",
        "You have 60 minutes to answer all five open-response questions in your own words; do not copy source material verbatim.",
        "Only the course slides, Learner Guide and activity resources permitted by the assessor may be used. Place phones and other materials under the table.",
        "Use a black or blue pen for hard-copy work. Do not use correction fluid or correction tape.",
        "Complete your answers on the document provided. Protect candidate, employee and organisational information.",
        f"The result is Competent (C) or Not Yet Competent (NYC). Competent requires sufficient evidence across {competency}.",
        "If assessed NYC, you may be offered re-assessment. You may appeal the assessment outcome through the stated ATO process.",
    ]
    for i,item in enumerate(items,1): doc.add_paragraph(f"{i}. {item}")
    p=doc.add_paragraph("8. Upload the completed answers to the LMS at "); add_hyperlink(p,LMS_URL,LMS_URL)
    doc.add_heading("Grading",1)
    table(doc,["Field","Assessor entry"],[("Grade","C / NYC (delete as appropriate)"),("Assessor Name",""),("Assessor NRIC",""),("Date",""),("Assessor Signature","")])
    page_break(doc)

def make_paper(kind,items,scenario=None):
    title="Written Assessment" if kind=="WA" else "Case Study Assessment"
    d=Document(); style_doc(d); assessment_header_footer(d,title); cover(d,title); instructions(d,title,"K1–K5" if kind=="WA" else "A1–A5")
    if scenario:
        d.add_heading("Case Study Scenario",1); d.add_paragraph(scenario); d.add_paragraph("Use this scenario for Questions 1–5 (A1–A5).")
        page_break(d)
    for i,(code,q,_,marks) in enumerate(items,1):
        d.add_heading(f"Question {i}: {code} ({marks} marks)",2); d.add_paragraph(q); answer_box(d,height_pt=90)
        if i%2==0 and i<len(items): page_break(d)
    d.add_heading("Candidate declaration",1); d.add_paragraph("I declare that this submission is my own work and that I have followed the assessment instructions.")
    d.add_paragraph("Candidate signature: ______________________________  Date: __________________")
    name=f"WA (SAQ) - {META['title']} - v{VER}.docx" if kind=="WA" else f"CS Assessment - {META['title']} - v{VER}.docx"
    out=ASSESSMENT_DIR/name; d.save(out); normalize_docx(out); return out

def make_key(kind,items,scenario=None):
    title=("Written Assessment" if kind=="WA" else "Case Study Assessment")+" Answer Key"
    d=Document(); style_doc(d); assessment_header_footer(d,"Confidential Answer Key"); cover(d,title,True)
    d.add_heading("Assessor guidance",1)
    d.add_paragraph("Use professional judgement. Accept equivalent technically sound responses that demonstrate the mapped evidence. Do not award full credit for unsupported lists where explanation, application or justification is required.")
    if scenario: d.add_heading("Case scenario reference",1); d.add_paragraph(scenario)
    for i,(code,q,ans,marks) in enumerate(items,1):
        d.add_heading(f"Question {i}: {code} — {marks} marks",2); d.add_paragraph(q)
        each=marks/len(ans); answer_box(d,[f"{x} [up to {each:g} marks]" for x in ans])
        if i<len(items): page_break(d)
    page_break(d); d.add_heading("Assessment decision record",1)
    table(d,["Question","Maximum","Awarded","Evidence / feedback"],[(x[0],x[3],"","") for x in items]+[("Total",100,"","")])
    d.add_paragraph("Final decision:  ☐ Competent   ☐ Not Yet Competent")
    d.add_paragraph("Assessor name/signature: __________________________________  Date: ______________")
    name=f"Answer to WA (SAQ) - {META['title']} - v{VER}.docx" if kind=="WA" else f"Answer to CS Assessment - {META['title']} - v{VER}.docx"
    out=ASSESSMENT_DIR/name; d.save(out); normalize_docx(out); return out

def main():
    ASSESSMENT_DIR.mkdir(parents=True,exist_ok=True)
    docs=[make_paper('WA',WA),make_key('WA',WA),make_paper('CS',CS,CASE_SCENARIO),make_key('CS',CS,CASE_SCENARIO)]
    print("\n".join(str(x) for x in docs))

if __name__=='__main__': main()
