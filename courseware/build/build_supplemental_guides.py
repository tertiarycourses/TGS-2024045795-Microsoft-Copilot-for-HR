"""Build answer-free facilitator and assessment plans for TGS-2024045795."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from course_data import COURSEWARE, LEARNING_OUTCOMES, META


BLUE = RGBColor(8, 76, 139)
INK = RGBColor(23, 42, 63)


def make_doc(kind: str, version: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = INK
    for name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 11)]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = BLUE

    p = doc.add_paragraph(style="Title")
    p.add_run(kind.upper())
    doc.add_paragraph(META["title"], style="Heading 1")
    doc.add_paragraph(f"Course Code: {META['code']}  |  {META['duration']}  |  {META['tsc']}")
    doc.add_paragraph(f"Tertiary Infotech Academy Pte Ltd  |  UEN 201200696W")
    doc.add_paragraph(f"Version {version}  |  {META['date']}")
    doc.add_paragraph("This document contains assessment administration and evidence requirements only. Candidate question papers and confidential answer keys are controlled separately.")

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{META['code']}  |  {kind}  |  Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
    doc.add_paragraph()


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def facilitator() -> Path:
    doc = make_doc("Facilitator Guide", "6.3")
    doc.add_heading("Delivery intent", 1)
    doc.add_paragraph("Use visual slides to frame decisions and the Learner Guide and numbered Labs for detailed procedures. Work with synthetic HR data and approved Microsoft 365 tenant content. Every consequential HR decision retains a named human owner.")
    doc.add_heading("Learning outcomes", 1)
    for i, outcome in enumerate(LEARNING_OUTCOMES, 1):
        doc.add_paragraph(f"LO{i}. {outcome}", style="List Number")
    doc.add_heading("Before the class", 1)
    bullets(doc, [
        "Confirm trainer and learner access to the current slides, Learner Guide, Labs and licensed Microsoft 365/Copilot tools; prepare a no-tenant fallback using the supplied mock data.",
        "Check that the 12 Lab folders contain instructions, prompts, synthetic data, workflow diagrams and evidence workbooks. Demonstrate one complete Test-it record before independent practice.",
        "Confirm assessment venue, timing, candidate identity process and accessible accommodations under the approved ATO procedure. Keep question papers controlled and answer keys out of learner links.",
    ])
    doc.add_heading("Day 1 run of show", 1)
    add_table(doc, ["Time", "Facilitation", "Evidence to collect"], [
        ["09:30–11:00", "Welcome, outcomes and Unit 1: From Generative AI to Agentic HR", "Distinguish rules, Copilot and bounded agents; name the human owner"],
        ["11:00–13:00", "Reflection, then parallel Labs 1–3: talent need, job description, screening", "Source citations, job-related criteria, reviewable evidence matrix"],
        ["14:00–15:30", "Unit 2: Microsoft 365 Copilot Across the HR Lifecycle", "Show a prompt, grounded response, verification and approval"],
        ["15:30–18:30", "Reflection, parallel Labs 4–6 and teach-back", "Structured interview, onboarding and development evidence"],
    ])
    doc.add_page_break()
    doc.add_heading("Day 2 run of show", 1)
    add_table(doc, ["Time", "Facilitation", "Evidence to collect"], [
        ["09:30–11:00", "Retrieval practice and Unit 3: Build Governed HR Agents", "Purpose, knowledge, instructions, tools, boundaries and tests"],
        ["11:00–13:00", "Reflection and parallel Labs 7–9: policy, leave and benefits agents", "Citation, permission and escalation tests; idempotent leave flow"],
        ["14:00–16:00", "Unit 4: Responsible Adoption; reflection and Labs 10–12", "Privacy, fairness, owner and scale/stop decision"],
        ["16:00–18:30", "Assessment briefing; 16:30–17:30 WA; 17:30–18:30 Case Study", "Individual candidate submissions and assessor records"],
    ])
    doc.add_paragraph("Lunch is 13:00–14:00 on both days. The two 15-minute facilitated reflections each day are contact time. Labs are parallel pods; each learner completes an assigned in-class lab and shares evidence. The remaining labs are self-paced extensions, as detailed in the Lesson Plan.")
    doc.add_heading("Lab coaching and control points", 1)
    add_table(doc, ["Labs", "Observe", "Intervene when"], [
        ["1–3", "Cited talent evidence, inclusive criteria, candidate correction route", "AI is allowed to decide or reject a candidate"],
        ["4–6", "Consistent interviews, approved onboarding tasks, employee-led development", "A model infers personality or commits employment terms"],
        ["7–9", "Current policy citation, authenticated personal status, deterministic approval", "An agent claims entitlement or writes twice without confirmation"],
        ["10–12", "Aggregated wellness data, controlled offboarding, balanced control-room metrics", "Individual health inference or access removal lacks approval"],
    ])
    doc.add_heading("Assessment conduct", 1)
    bullets(doc, [
        "Use the current Written Assessment and Case Study candidate papers, version 3.2. Each is 60 minutes, individual and open-book under the conditions printed on the paper.",
        "Allow only the materials named in the candidate instructions. Explain C/NYC decisions, the appeal route and the re-assessment procedure without coaching answers.",
        "Protect submissions and assessor records. Put confidential answer keys in trainer-controlled storage; never attach them to TMS courseware links or public repositories.",
    ])
    doc.add_heading("After delivery", 1)
    bullets(doc, [
        "Check that each learner's lab evidence names the source, the human reviewer, exceptions and one Test-it result.",
        "Record attendance, individual assessment decisions, feedback and follow-up in the approved TMS workflow.",
        "Escalate product access, licence or connector limitations as environment constraints; use synthetic fallback evidence and mark any unperformed live test clearly.",
    ])
    path = COURSEWARE / "FG-Microsoft Copilot for HR.docx"
    doc.save(path)
    return path


def assessment_plan() -> Path:
    doc = make_doc("Assessment Plan", "4.1")
    doc.add_heading("Assessment purpose and scope", 1)
    doc.add_paragraph("Assess the candidate's knowledge and applied ability for Human Resource Digitalisation-4 (HRS-HRM-4031-1.1) using the current Microsoft Copilot for HR course outcomes. The assessment uses one Written Assessment (WA-SAQ) and one Case Study (CS); no Practical Performance paper is selected for this release.")
    doc.add_heading("Assessment methods and conditions", 1)
    add_table(doc, ["Instrument", "Duration", "Evidence", "Decision"], [
        ["WA-SAQ v3.2", "60 minutes", "Five individual open-response items, K1–K5; 20 marks each", "Apply the controlled assessor rubric to each knowledge item"],
        ["Case Study v3.2", "60 minutes", "HarbourLight Logistics scenario and five individual open-response items, A1–A5; 20 marks each", "Apply the controlled assessor rubric to the scenario evidence"],
    ])
    doc.add_paragraph("Both papers are individual, open-book assessments. Candidate identity, permitted materials, submission method, conduct rules and C/NYC outcome language are stated on the papers. The assessor records sufficiency against every required item; a total score alone does not replace the competency decision.")
    doc.add_page_break()
    doc.add_heading("Evidence map", 1)
    add_table(doc, ["Item", "Evidence requested", "Course outcome"], [
        ["K1", "Compare automation, Microsoft 365 Copilot and agents with HR examples and human roles", "LO1"],
        ["K2", "Explain lifecycle use cases and verification before use", "LO2"],
        ["K3", "Explain purpose, knowledge, instructions, tools, boundaries and tests", "LO2, LO3"],
        ["K4", "Explain responsible-AI controls and the risks they reduce", "LO4"],
        ["K5", "Outline a 90-day pilot with baseline, measures, ownership and scale/stop gate", "LO3, LO4"],
        ["A1", "Select proportionate capability types for scenario HR work", "LO1, LO2"],
        ["A2", "Design a fair and reviewable screening workflow", "LO2, LO4"],
        ["A3", "Design a controlled leave agent with validation, approval and audit", "LO2, LO3, LO4"],
        ["A4", "Contain and remediate obsolete policy citations", "LO3, LO4"],
        ["A5", "Recommend scale, remediation, pause or retirement using balanced measures", "LO3, LO4"],
    ])
    doc.add_heading("Competency decision and quality controls", 1)
    bullets(doc, [
        "Use the separately controlled marking rubric and suggested answers available only to authorised assessors. This plan contains no model responses or answer keys.",
        "Check evidence for validity, authenticity, sufficiency and currency. Record a reason when evidence is insufficient, contradictory or copied without application to the scenario.",
        "Make the C/NYC decision using evidence across K1–K5 and A1–A5. Complete candidate and assessor records, provide feedback, and follow the approved re-assessment and appeal process where required.",
        "Handle personal data and submissions through the approved TMS and assessment-record locations; restrict access by role and retain records under the ATO policy.",
    ])
    doc.add_heading("Assessment administration sequence", 1)
    add_table(doc, ["Stage", "Assessor action"], [
        ["Before", "Confirm candidate identity, accommodations, assessment versions, venue and allowed materials; explain timing, C/NYC, re-assessment and appeal."],
        ["During", "Supervise the individual 60-minute WA followed by the 60-minute CS; answer process questions only and record incidents."],
        ["After", "Secure submissions, apply the controlled rubric independently, record evidence and feedback, complete required signatures and TMS status."],
    ])
    doc.add_heading("Version and distribution", 1)
    doc.add_paragraph("This version 4.1 plan aligns with the v6.3 courseware and v3.2 candidate papers dated 27 September 2026. The superseded v4.0 plan contains old course content and answer annexes and must not be used as the current TMS link. Confidential marking material is maintained separately.")
    path = COURSEWARE / "AP-Microsoft Copilot for HR.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for generated in [facilitator(), assessment_plan()]:
        print(generated)
